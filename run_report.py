"""Run 'Consolidated Employee Tips and Hours' for a date range and read the result.

The date boxes are readonly - typing is refused -
so the value is set on the Angular model behind the picker instead. Read-only: it
runs a report, it changes nothing.
"""
import sys, json
from playwright.sync_api import sync_playwright

import settings
BASE = settings.BASE


def run(start_iso, end_iso, dump=None):
    calls = []
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        ctx = b.new_context(viewport={"width": 1280, "height": 720}, storage_state="session.json")
        pg = ctx.new_page()
        pg.on("request", lambda r: calls.append(r.method + " " + r.url)
              if r.resource_type in ("xhr", "fetch") and "report" in r.url.lower() else None)
        pg.goto(BASE + "/app.html#/reporting/report/9000002", wait_until="load", timeout=60000)
        pg.wait_for_timeout(14000)
        pg.get_by_text("Select All", exact=True).first.click()
        pg.wait_for_timeout(1500)

        set_ok = pg.evaluate("""([s, e]) => {
            const pick = [...document.querySelectorAll('md-datepicker')];
            const out = [];
            pick.forEach(el => {
                const name = el.getAttribute('name');
                const sc = angular.element(el).scope();
                if (!sc || !sc.psc || !sc.psc.parameterObject) { out.push(name + ':no-scope'); return; }
                const iso = name === 'FROM_DATE' ? s : e;
                const [y, m, d] = iso.split('-').map(Number);
                sc.psc.parameterObject.value = new Date(y, m - 1, d);
                sc.$apply();
                out.push(name + ':set');
            });
            return out;
        }""", [start_iso, end_iso])
        pg.wait_for_timeout(2000)
        shown = [x.input_value() for x in pg.query_selector_all("input#datepicker-syn-input-id")]
        print("model set:", set_ok, "| boxes now show:", shown)
        if not all(v for v in shown):
            b.close()
            raise SystemExit("dates did not take - refusing to run")

        pg.get_by_role("button", name="Run").first.click()
        pg.wait_for_timeout(35000)
        if "Download Report" not in pg.inner_text("body"):
            b.close()
            raise SystemExit("report did not finish - no Download Report button")

        # The rendered report lives in an iframe, so take the xlsx rather than
        # scraping the screen. Locate the menu item by text: the menu re-renders
        # at a different scale once the report loads and old coordinates land on
        # "PDF Report" instead.
        pg.get_by_text("Download Report", exact=False).first.click()
        pg.wait_for_timeout(2500)
        with pg.expect_download(timeout=120000) as dl:
            pg.get_by_text("Excel Report", exact=False).first.click()
        path = dump or "report.xlsx"
        dl.value.save_as(path)
        b.close()
    return path


if __name__ == "__main__":
    s, e = sys.argv[1], sys.argv[2]
    path = run(s, e, dump="hours_%s_%s.xlsx" % (s, e))
    print("saved", path)
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    store = None
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 13)]
        joined = " ".join(str(v) for v in vals if v is not None)
        if "Store" in joined:
            store = joined.strip(); print(store); continue
        name = ws.cell(r, 2).value
        if name and ws.cell(r, 7).value is not None:
            print("   %-24s sect=%-14s rate=%-6s reg=%-7s ot=%s"
                  % (name, ws.cell(r, 4).value, ws.cell(r, 6).value,
                     ws.cell(r, 7).value, ws.cell(r, 8).value))
